"""Validaciones de estructura OOXML y fórmulas del reporte Excel."""

from __future__ import annotations

from datetime import UTC, datetime
from decimal import Decimal
from io import BytesIO
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import load_workbook

from agente_comprobantes.control_financiero.modelos import (
    OperacionFlete,
    RevisionExpediente,
    SituacionControl,
)
from agente_comprobantes.control_financiero.presentacion import construir_tablero
from agente_comprobantes.control_financiero.reporte_ejecutivo import (
    exportar_excel_ejecutivo,
)
from agente_comprobantes.control_financiero.repositorio_json import RepositorioControlFletes
from agente_comprobantes.dominio.enumeraciones import Moneda

NS = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


def op(
    ident: str, numero: str, monto: str, *,
    moneda: Moneda = Moneda.PEN, es_flete: bool | None = True,
) -> OperacionFlete:
    return OperacionFlete(
        identificador=ident,
        numero_factura=numero,
        proveedor="Empresa prueba",
        moneda=moneda,
        monto_facturado=Decimal(monto),
        monto_cotizado=Decimal("100"),
        codigo_cotizacion="COT-1" if moneda == Moneda.PEN else "COT-USD",
        es_flete=es_flete,
    )


def sheet(data: bytes, number: int) -> ET.Element:
    with ZipFile(BytesIO(data)) as z:
        return ET.fromstring(z.read(f"xl/worksheets/sheet{number}.xml"))


def valor(raiz: ET.Element, celda: str) -> str | None:
    obj = raiz.find(f".//x:c[@r='{celda}']", NS)
    if obj is None:
        return None
    v = obj.find("x:v", NS)
    if v is not None and v.text is not None:
        return v.text
    t = obj.find("x:is/x:t", NS)
    if t is not None:
        return t.text
    f = obj.find("x:f", NS)
    return f"={f.text}" if f is not None else None


def test_reporte_abre_como_paquete_xlsx_y_contiene_5_hojas() -> None:
    data = exportar_excel_ejecutivo(construir_tablero([]))
    # Abrir el archivo con el lector empleado por el propio proyecto
    # evita los falsos positivos de comprobar únicamente que el ZIP sea válido.
    libro = load_workbook(BytesIO(data), data_only=False)
    assert libro.sheetnames == [
        "Resumen ejecutivo", "Operaciones", "Observaciones",
        "Historial", "Leer primero",
    ]
    assert libro["Resumen ejecutivo"].freeze_panes == "A6"
    assert libro["Operaciones"]["A5"].value == "ID expediente"
    with ZipFile(BytesIO(data)) as zipfile:
        assert zipfile.testzip() is None
        assert "[Content_Types].xml" in zipfile.namelist()
        assert "xl/styles.xml" in zipfile.namelist()
        workbook = ET.fromstring(zipfile.read("xl/workbook.xml"))
        names = [
            el.attrib["name"] for el in workbook.findall("x:sheets/x:sheet", NS)
        ]
        assert names == [
            "Resumen ejecutivo", "Operaciones", "Observaciones",
            "Historial", "Leer primero",
        ]


def test_importes_y_formulas_recalculables_en_excel() -> None:
    a = op("FAC-1", "F001-1", "80")
    b = op("FAC-2", "F001-2", "45")
    libro = exportar_excel_ejecutivo(construir_tablero([a, b]))
    resumen, operaciones = sheet(libro, 1), sheet(libro, 2)
    assert valor(resumen, "B6") == "100"
    assert valor(resumen, "C6") == "125"
    assert valor(resumen, "D6") == "=C6-B6"
    assert valor(operaciones, "K6") == "=I6-J6"
    assert valor(operaciones, "P6") == "=K6-O6"
    formula = operaciones.find(".//x:c[@r='P6']/x:f", NS)
    assert formula is not None and formula.text == "K6-O6"


def test_presupuesto_no_incluye_materiales_ni_mezcla_monedas() -> None:
    a = op("FAC-1", "F001-1", "130")
    b = op("FAC-2", "F001-2", "900", es_flete=False)
    c = op("FAC-3", "F001-3", "20", moneda=Moneda.USD)
    libro = exportar_excel_ejecutivo(construir_tablero([a, b, c]))
    resumen = sheet(libro, 1)
    assert valor(resumen, "A6") == "PEN"
    assert valor(resumen, "B6") == "100"
    assert valor(resumen, "A7") == "USD"
    assert valor(resumen, "B7") == "100"
    assert len(sheet(libro, 2).findall("x:sheetData/x:row", NS)) == 6


def test_filtrado_del_detalle_no_trunca_resumen_de_presupuestos() -> None:
    data = exportar_excel_ejecutivo(
        construir_tablero([
            op("FAC-1", "F001-1", "80"), op("FAC-2", "F001-2", "45"),
        ]),
        moneda=Moneda.PEN,
        situacion=SituacionControl.PENDIENTE,
        busqueda="F001-1",
    )
    assert valor(sheet(data, 1), "C6") == "125"
    assert valor(sheet(data, 2), "A6") == "FAC-1"
    assert valor(sheet(data, 2), "A7") is None


def test_texto_sospechoso_no_se_convierte_en_formula() -> None:
    original = op("FAC-1", "F001-1", "100")
    original = original.model_copy(update={"proveedor": "=HYPERLINK(\"http://ejemplo\")"})
    original.revisiones.append(
        RevisionExpediente(
            fecha=datetime(2026, 10, 26, tzinfo=UTC),
            motivo="@SUM(1,1)", cambios={"guia": "+cmd"},
        )
    )
    data = exportar_excel_ejecutivo(construir_tablero([original]))
    operaciones = sheet(data, 2)
    assert valor(operaciones, "D6").startswith("=HYPERLINK")
    assert operaciones.find(".//x:c[@r='D6']/x:f", NS) is None
    historial = sheet(data, 4)
    assert valor(historial, "C6") == "@SUM(1,1)"
    assert historial.find(".//x:c[@r='C6']/x:f", NS) is None
    recargado = load_workbook(BytesIO(data))
    assert recargado["Operaciones"]["D6"].data_type == "s"
    assert recargado["Historial"]["C6"].data_type == "s"


def test_generacion_no_modifica_json_o_excel(tmp_path) -> None:
    ruta = tmp_path / "control_fletes.json"
    excel = tmp_path / "registros.xlsx"
    excel.write_bytes(b"copia-de-prueba")
    repo = RepositorioControlFletes(ruta)
    repo.agregar(op("FAC-1", "F001-1", "100"))
    json_anterior = ruta.read_bytes()
    generado = exportar_excel_ejecutivo(construir_tablero(repo.listar()))
    assert generado.startswith(b"PK")
    assert ruta.read_bytes() == json_anterior
    assert excel.read_bytes() == b"copia-de-prueba"



def test_el_archivo_excel_se_abre_con_datos_y_formulas() -> None:
    data = exportar_excel_ejecutivo(
        construir_tablero([op("FAC-1", "F001-1", "80")])
    )
    archivo = load_workbook(BytesIO(data), data_only=False)
    resumen = archivo["Resumen ejecutivo"]
    detalle = archivo["Operaciones"]
    assert resumen["B6"].value == 100
    assert detalle["A6"].value == "FAC-1"
    assert detalle["K6"].value == "=I6-J6"
    assert detalle["O6"].value == "=M6+N6"
    assert detalle["P6"].value == "=K6-O6"
    assert detalle.auto_filter.ref.startswith("A5:")
    assert detalle["A1"].font.bold
