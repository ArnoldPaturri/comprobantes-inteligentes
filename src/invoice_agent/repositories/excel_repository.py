from __future__ import annotations

import os
import re
import tempfile
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from invoice_agent.domain import (
    Currency,
    DocumentReference,
    DocumentType,
    InvoiceRecord,
    PaymentLinkStatus,
    PaymentReconciliationResult,
    PaymentRecord,
    PaymentStatus,
    QuotationRecord,
    QuotationStatus,
    ReconciliationResult,
    VerificationStatus,
)
from invoice_agent.exceptions import DuplicateCodeError, RepositoryError
from invoice_agent.repositories.base import DocumentRepository

ILLEGAL_EXCEL_CHARACTER_PATTERN = re.compile(
    r"[\x00-\x08\x0B-\x0C\x0E-\x1F]"
)

QUOTATION_HEADERS = [
    "Cód. Cotización",
    "Proveedor",
    "RUC Proveedor",
    "Concepto / Descripción",
    "Monto Cotizado",
    "Moneda",
    "Fecha Cotización",
    "Estado",
    "Estado Verificación",
    "Hash Documento",
    "Archivos Origen",
    "Campos Inciertos",
    "Fecha Registro",
    "Acción / Observaciones",
]

INVOICE_HEADERS = [
    "ID Factura",
    "Cód. Cotización",
    "Proveedor",
    "RUC Proveedor",
    "N° Factura",
    "Concepto / Descripción",
    "Monto Facturado",
    "Moneda",
    "Fecha Factura",
    "Estado Verificación",
    "Estado de Pago",
    "Hash Documento",
    "Archivos Origen",
    "Campos Inciertos",
    "Posible Duplicado De",
    "Fecha Registro",
    "Acción / Observaciones",
]

COMPARISON_HEADERS = [
    "Cód. Cotización",
    "Proveedor",
    "Concepto / Descripción",
    "N° Factura",
    "Monto Cotizado",
    "Monto Facturado",
    "Diferencia ($)",
    "% Desvío",
    "Diagnóstico",
    "Estado Verificación",
    "Estado de Pago",
    "Fecha Factura",
    "Acción / Observaciones",
]

PAYMENT_HEADERS = [
    "ID Pago",
    "ID Factura",
    "Tipo de Pago",
    "Entidad",
    "N° Operación",
    "Monto Pagado",
    "Moneda",
    "Fecha Pago",
    "Estado Verificación",
    "Hash Documento",
    "Fecha Registro",
    "Acción / Observaciones",
    "Cód. Cotización",
    "Estado Conciliación",
    "Archivos Origen",
    "Posible Duplicado De",
]

SHEETS = {
    "COTIZACIONES": QUOTATION_HEADERS,
    "FACTURAS": INVOICE_HEADERS,
    "COMPARATIVA": COMPARISON_HEADERS,
    "PAGOS": PAYMENT_HEADERS,
}


class ExcelRepository(DocumentRepository):
    """Repositorio XLSX con guardado atómico para reducir riesgo de corrupción."""

    def __init__(self, path: str | Path) -> None:
        self._path = Path(path).expanduser().resolve()

    @property
    def location(self) -> str:
        return str(self._path)

    def initialize(self) -> None:
        self._path.parent.mkdir(parents=True, exist_ok=True)
        if self._path.exists():
            workbook = load_workbook(self._path)
            changed = False
            for name, headers in SHEETS.items():
                if name not in workbook.sheetnames:
                    sheet = workbook.create_sheet(name)
                    self._initialize_sheet(sheet, headers)
                    changed = True
                else:
                    changed = self._ensure_headers(workbook[name], headers) or changed
            if changed:
                self._save_atomic(workbook)
            else:
                workbook.close()
            return

        workbook = Workbook()
        default = workbook.active
        workbook.remove(default)
        for name, headers in SHEETS.items():
            sheet = workbook.create_sheet(name)
            self._initialize_sheet(sheet, headers)
        self._save_atomic(workbook)

    def get_quotation(self, quotation_code: str) -> QuotationRecord | None:
        normalized = quotation_code.strip().upper()
        for row in self.read_sheet("COTIZACIONES"):
            if str(row["Cód. Cotización"]).strip().upper() == normalized:
                return self._quotation_from_row(row)
        return None

    def list_quotations(self) -> list[QuotationRecord]:
        self.initialize()

        quotations = [
            self._quotation_from_row(row)
            for row in self.read_sheet("COTIZACIONES")
        ]

        return sorted(
            quotations,
            key=lambda quotation: quotation.registered_at,
            reverse=True,
        )

    def add_quotation(self, quotation: QuotationRecord) -> None:
        self.initialize()
        if self.get_quotation(quotation.quotation_code):
            raise DuplicateCodeError(
                f"Ya existe la cotización {quotation.quotation_code}; no se creó otra fila."
            )

        workbook = load_workbook(self._path)
        sheet = workbook["COTIZACIONES"]
        self._append_mapping(
            sheet,
            QUOTATION_HEADERS,
            {
                "Cód. Cotización": quotation.quotation_code,
                "Proveedor": quotation.provider,
                "RUC Proveedor": quotation.provider_tax_id,
                "Concepto / Descripción": quotation.concept,
                "Monto Cotizado": quotation.quoted_amount,
                "Moneda": quotation.currency,
                "Fecha Cotización": quotation.quotation_date,
                "Estado": quotation.status,
                "Estado Verificación": quotation.verification_status,
                "Hash Documento": quotation.source_hash,
                "Archivos Origen": quotation.source_files,
                "Campos Inciertos": quotation.uncertain_fields,
                "Fecha Registro": quotation.registered_at,
                "Acción / Observaciones": quotation.observations,
                
            },
        )
        self._save_atomic(workbook)

    def add_invoice(
        self,
        invoice: InvoiceRecord,
        quotation: QuotationRecord,
        reconciliation: ReconciliationResult,
    ) -> None:
        self.initialize()
        workbook = load_workbook(self._path)
        self._append_mapping(
            workbook["FACTURAS"],
            INVOICE_HEADERS,
            {
                "ID Factura": invoice.invoice_id,
                "Cód. Cotización": invoice.quotation_code,
                "Proveedor": invoice.provider,
                "RUC Proveedor": invoice.provider_tax_id,
                "N° Factura": invoice.invoice_number,
                "Concepto / Descripción": invoice.concept,
                "Monto Facturado": invoice.invoiced_amount,
                "Moneda": invoice.currency,
                "Fecha Factura": invoice.invoice_date,
                "Estado Verificación": invoice.verification_status,
                "Estado de Pago": invoice.payment_status,
                "Hash Documento": invoice.source_hash,
                "Archivos Origen": invoice.source_files,
                "Campos Inciertos": invoice.uncertain_fields,
                "Posible Duplicado De": invoice.duplicate_of,
                "Fecha Registro": invoice.registered_at,
                "Acción / Observaciones": invoice.observations,
            },
        )
        self._append_mapping(
            workbook["COMPARATIVA"],
            COMPARISON_HEADERS,
            {
                "Cód. Cotización": invoice.quotation_code,
                "Proveedor": invoice.provider or quotation.provider,
                "Concepto / Descripción": invoice.concept or quotation.concept,
                "N° Factura": invoice.invoice_number,
                "Monto Cotizado": reconciliation.quoted_amount,
                "Monto Facturado": reconciliation.invoiced_amount,
                "Diferencia ($)": reconciliation.difference,
                "% Desvío": reconciliation.deviation_percent,
                "Diagnóstico": reconciliation.diagnosis,
                "Estado Verificación": invoice.verification_status,
                "Estado de Pago": invoice.payment_status,
                "Fecha Factura": invoice.invoice_date,
                "Acción / Observaciones": invoice.observations,
            },
        )
        self._update_quotation_status(
            workbook["COTIZACIONES"], invoice.quotation_code, QuotationStatus.INVOICED
        )
        self._save_atomic(workbook)

    def get_invoice(self, invoice_id: str) -> InvoiceRecord | None:
        normalized = invoice_id.strip().upper()
        for row in self.read_sheet("FACTURAS"):
            if str(row["ID Factura"]).strip().upper() == normalized:
                return self._invoice_from_row(row)
        return None

    def list_invoices(
        self,
        quotation_code: str | None = None,
    ) -> list[InvoiceRecord]:
        self.initialize()

        invoices = [
            self._invoice_from_row(row)
            for row in self.read_sheet("FACTURAS")
        ]

        if quotation_code is not None:
            normalized_code = quotation_code.strip().upper()

            invoices = [
                invoice
                for invoice in invoices
                if invoice.quotation_code.strip().upper()
                == normalized_code
            ]

        return sorted(
            invoices,
            key=lambda invoice: invoice.registered_at,
            reverse=True,
        )

    def get_verified_payment_total(self, invoice_id: str) -> Decimal:
        normalized = invoice_id.strip().upper()
        total = Decimal("0")
        for row in self.read_sheet("PAGOS"):
            if str(row["ID Factura"]).strip().upper() != normalized:
                continue
            if str(row["Estado Verificación"]) != VerificationStatus.VERIFIED.value:
                continue
            amount = row["Monto Pagado"]
            if amount not in (None, ""):
                total += Decimal(str(amount))
        return total

    def get_payment(self, payment_id: str) -> PaymentRecord | None:
        normalized = payment_id.strip().upper()

        for row in self.read_sheet("PAGOS"):
            if str(row["ID Pago"]).strip().upper() == normalized:
                return self._payment_from_row(row)

        return None

    def list_pending_payments(
        self,
        quotation_code: str,
    ) -> list[PaymentRecord]:
        normalized = quotation_code.strip().upper()
        pending: list[PaymentRecord] = []

        for row in self.read_sheet("PAGOS"):
            row_quotation = str(
                row["Cód. Cotización"] or ""
            ).strip().upper()

            if row_quotation != normalized:
                continue

            if (
                str(row["Estado Conciliación"] or "")
                != PaymentLinkStatus.PENDING_INVOICE.value
            ):
                continue

            pending.append(self._payment_from_row(row))

        return pending
    
    def add_payment(
        self,
        payment: PaymentRecord,
        invoice: InvoiceRecord | None = None,
        reconciliation: PaymentReconciliationResult | None = None,
    ) -> None:
        self.initialize()
        workbook = load_workbook(self._path)
        self._append_mapping(
            workbook["PAGOS"],
            PAYMENT_HEADERS,
            {
                "ID Pago": payment.payment_id,
                "ID Factura": payment.invoice_id,
                "Tipo de Pago": payment.payment_type,
                "Entidad": payment.payment_entity,
                "N° Operación": payment.operation_number,
                "Monto Pagado": payment.paid_amount,
                "Moneda": payment.currency,
                "Fecha Pago": payment.payment_date,
                "Estado Verificación": payment.verification_status,
                "Hash Documento": payment.source_hash,
                "Fecha Registro": payment.registered_at,
                "Acción / Observaciones": payment.observations,
                "Cód. Cotización": payment.quotation_code,
                "Estado Conciliación": payment.reconciliation_status,
                "Archivos Origen": payment.source_files,
                "Posible Duplicado De": payment.duplicate_of,
            },
        )

        if (
            invoice is not None
            and reconciliation is not None
            and reconciliation.applied
        ):
            self._update_invoice_payment_status(
                workbook["FACTURAS"],
                invoice.invoice_id,
                reconciliation.resulting_payment_status,
            )
            self._update_comparison_payment_status(
                workbook["COMPARATIVA"],
                invoice.quotation_code,
                invoice.invoice_number,
                reconciliation.resulting_payment_status,
            )
        self._save_atomic(workbook)

    def link_payment(
        self,
        payment: PaymentRecord,
        invoice: InvoiceRecord,
        reconciliation: PaymentReconciliationResult,
    ) -> None:
        self.initialize()
        workbook = load_workbook(self._path)
        sheet = workbook["PAGOS"]

        headers = {
            cell.value: cell.column
            for cell in sheet[1]
        }

        for row_number in range(2, sheet.max_row + 1):
            current = sheet.cell(
                row=row_number,
                column=headers["ID Pago"],
            ).value

            if (
                str(current).strip().upper()
                != payment.payment_id.strip().upper()
            ):
                continue

            sheet.cell(
                row=row_number,
                column=headers["ID Factura"],
            ).value = payment.invoice_id

            sheet.cell(
                row=row_number,
                column=headers["Estado Conciliación"],
            ).value = payment.reconciliation_status.value

            sheet.cell(
                row=row_number,
                column=headers["Acción / Observaciones"],
            ).value = self._excel_value(payment.observations)

            break
        else:
            workbook.close()
            raise RepositoryError(
                f"No existe el pago {payment.payment_id}."
            )

        if reconciliation.applied:
            self._update_invoice_payment_status(
                workbook["FACTURAS"],
                invoice.invoice_id,
                reconciliation.resulting_payment_status,
            )

            self._update_comparison_payment_status(
                workbook["COMPARATIVA"],
                invoice.quotation_code,
                invoice.invoice_number,
                reconciliation.resulting_payment_status,
            )

        self._save_atomic(workbook)

    def list_document_references(self) -> list[DocumentReference]:
        self.initialize()
        references: list[DocumentReference] = []
        for row in self.read_sheet("COTIZACIONES"):
            references.append(
                DocumentReference(
                    document_id=str(row["Cód. Cotización"]),
                    document_type=DocumentType.QUOTATION,
                    provider=self._optional_str(row["Proveedor"]),
                    provider_tax_id=self._optional_str(row["RUC Proveedor"]),
                    concept=self._optional_str(row["Concepto / Descripción"]),
                    document_date=row["Fecha Cotización"],
                    amount=row["Monto Cotizado"],
                    currency=self._optional_currency(row["Moneda"]),
                    source_hash=str(row["Hash Documento"]),
                )
            )
        for row in self.read_sheet("FACTURAS"):
            references.append(
                DocumentReference(
                    document_id=str(row["ID Factura"]),
                    document_type=DocumentType.INVOICE,
                    provider=self._optional_str(row["Proveedor"]),
                    provider_tax_id=self._optional_str(row["RUC Proveedor"]),
                    concept=self._optional_str(row["Concepto / Descripción"]),
                    document_number=self._optional_str(row["N° Factura"]),
                    document_date=row["Fecha Factura"],
                    amount=row["Monto Facturado"],
                    currency=self._optional_currency(row["Moneda"]),
                    source_hash=str(row["Hash Documento"]),
                )
            )
        for row in self.read_sheet("PAGOS"):
            payment_type = self._optional_document_type(row["Tipo de Pago"])
            if payment_type is None:
                continue
            references.append(
                DocumentReference(
                    document_id=str(row["ID Pago"]),
                    document_type=payment_type,
                    document_date=row["Fecha Pago"],
                    amount=row["Monto Pagado"],
                    currency=self._optional_currency(row["Moneda"]),
                    operation_number=self._optional_str(row["N° Operación"]),
                    source_hash=str(row["Hash Documento"]),
                )
            )
        return references

    def read_sheet(self, sheet_name: str) -> list[dict[str, object]]:
        if not self._path.exists():
            return []
        workbook = load_workbook(self._path, data_only=True, read_only=True)
        try:
            if sheet_name not in workbook.sheetnames:
                raise RepositoryError(f"No existe la hoja {sheet_name}.")
            sheet = workbook[sheet_name]
            rows = sheet.iter_rows(values_only=True)
            headers = next(rows, None)
            if not headers:
                return []
            return [dict(zip(headers, row, strict=False)) for row in rows if any(row)]
        finally:
            workbook.close()

    @staticmethod
    def _initialize_sheet(sheet: Any, headers: list[str]) -> None:
        sheet.append(headers)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
        fill = PatternFill("solid", fgColor="1F4E78")
        for cell in sheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for index, header in enumerate(headers, start=1):
            width = min(max(len(header) + 3, 14), 32)
            sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.row_dimensions[1].height = 28

    @staticmethod
    def _ensure_headers(sheet: Any, expected: list[str]) -> bool:
        actual = [cell.value for cell in sheet[1]]

        while actual and actual[-1] is None:
            actual.pop()

        if actual == expected:
            return False

        if actual != expected[: len(actual)]:
            raise RepositoryError(
                f"La hoja {sheet.title} existe, pero sus columnas "
                "no coinciden con el contrato."
            )

        fill = PatternFill("solid", fgColor="1F4E78")

        for index, header in enumerate(
            expected[len(actual):],
            start=len(actual) + 1,
        ):
            cell = sheet.cell(row=1, column=index, value=header)
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = fill
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )
            sheet.column_dimensions[get_column_letter(index)].width = min(
                max(len(header) + 3, 14),
                32,
            )

        sheet.auto_filter.ref = (
            f"A1:{get_column_letter(len(expected))}{sheet.max_row}"
        )
        return True

    @classmethod
    def _append_mapping(cls, sheet: Any, headers: list[str], mapping: dict[str, Any]) -> None:
        sheet.append([cls._excel_value(mapping.get(header)) for header in headers])
        row_number = sheet.max_row
        for cell in sheet[row_number]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        cls._apply_number_formats(sheet, row_number, headers)
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{sheet.max_row}"

    @staticmethod
    def _apply_number_formats(sheet: Any, row_number: int, headers: list[str]) -> None:
        for index, header in enumerate(headers, start=1):
            cell = sheet.cell(row=row_number, column=index)
            if header.startswith("Monto") or header == "Diferencia ($)":
                cell.number_format = '#,##0.00;[Red]-#,##0.00'
            elif header == "% Desvío":
                cell.number_format = '0.0000;[Red]-0.0000'
            elif header.startswith("Fecha") and isinstance(cell.value, datetime):
                cell.number_format = "dd/mm/yyyy hh:mm"
            elif header.startswith("Fecha") and isinstance(cell.value, date):
                cell.number_format = "dd/mm/yyyy"

    @staticmethod
    def _excel_value(value: Any) -> Any:
        if isinstance(value, Decimal):
            return float(value)

        if hasattr(value, "value"):
            value = value.value

        if isinstance(value, list):
            value = " | ".join(
                str(item)
                for item in value
            )

        if isinstance(value, str):
            return ILLEGAL_EXCEL_CHARACTER_PATTERN.sub(
                "",
                value,
            )

        return value

    @staticmethod
    def _update_quotation_status(sheet: Any, code: str, status: QuotationStatus) -> None:
        headers = {cell.value: cell.column for cell in sheet[1]}
        code_col = headers["Cód. Cotización"]
        status_col = headers["Estado"]
        for row_number in range(2, sheet.max_row + 1):
            current = sheet.cell(row=row_number, column=code_col).value
            if str(current).strip().upper() == code.strip().upper():
                sheet.cell(row=row_number, column=status_col).value = status.value
                return

    @staticmethod
    def _update_invoice_payment_status(
        sheet: Any,
        invoice_id: str,
        status: PaymentStatus,
    ) -> None:
        headers = {cell.value: cell.column for cell in sheet[1]}
        id_col = headers["ID Factura"]
        status_col = headers["Estado de Pago"]
        for row_number in range(2, sheet.max_row + 1):
            current = sheet.cell(row=row_number, column=id_col).value
            if str(current).strip().upper() == invoice_id.strip().upper():
                sheet.cell(row=row_number, column=status_col).value = status.value
                return

    @staticmethod
    def _update_comparison_payment_status(
        sheet: Any,
        quotation_code: str,
        invoice_number: str | None,
        status: PaymentStatus,
    ) -> None:
        headers = {cell.value: cell.column for cell in sheet[1]}
        quotation_col = headers["Cód. Cotización"]
        invoice_col = headers["N° Factura"]
        status_col = headers["Estado de Pago"]
        for row_number in range(2, sheet.max_row + 1):
            current_quotation = sheet.cell(row=row_number, column=quotation_col).value
            current_invoice = sheet.cell(row=row_number, column=invoice_col).value
            if (
                str(current_quotation).strip().upper() == quotation_code.strip().upper()
                and str(current_invoice or "").strip().upper()
                == str(invoice_number or "").strip().upper()
            ):
                sheet.cell(row=row_number, column=status_col).value = status.value
                return

    def _save_atomic(self, workbook: Any) -> None:
        self._path.parent.mkdir(parents=True, exist_ok=True)
        temporary_path: Path | None = None
        try:
            with tempfile.NamedTemporaryFile(
                dir=self._path.parent, prefix=f".{self._path.stem}-", suffix=".xlsx", delete=False
            ) as temporary:
                temporary_path = Path(temporary.name)
            workbook.save(temporary_path)
            workbook.close()
            os.replace(temporary_path, self._path)
        except Exception as exc:
            workbook.close()
            if temporary_path and temporary_path.exists():
                temporary_path.unlink()
            raise RepositoryError(f"No se pudo guardar el Excel: {self._path}") from exc

    @staticmethod
    def _quotation_from_row(row: dict[str, object]) -> QuotationRecord:
        return QuotationRecord(
            quotation_code=str(row["Cód. Cotización"]),
            provider=ExcelRepository._optional_str(row["Proveedor"]),
            provider_tax_id=ExcelRepository._optional_str(row["RUC Proveedor"]),
            concept=ExcelRepository._optional_str(row["Concepto / Descripción"]),
            quoted_amount=row["Monto Cotizado"],
            currency=ExcelRepository._optional_currency(row["Moneda"]),
            quotation_date=row["Fecha Cotización"],
            status=QuotationStatus(str(row["Estado"])),
            verification_status=VerificationStatus(str(row["Estado Verificación"])),
            source_hash=str(row["Hash Documento"]),
            source_files=ExcelRepository._split(row["Archivos Origen"]),
            uncertain_fields=ExcelRepository._split(row["Campos Inciertos"]),
            registered_at=row["Fecha Registro"] or datetime.now(),
            observations=ExcelRepository._split(row["Acción / Observaciones"]),
        )

    @staticmethod
    def _invoice_from_row(row: dict[str, object]) -> InvoiceRecord:
        return InvoiceRecord(
            invoice_id=str(row["ID Factura"]),
            quotation_code=str(row["Cód. Cotización"]),
            provider=ExcelRepository._optional_str(row["Proveedor"]),
            provider_tax_id=ExcelRepository._optional_str(row["RUC Proveedor"]),
            invoice_number=ExcelRepository._optional_str(row["N° Factura"]),
            concept=ExcelRepository._optional_str(row["Concepto / Descripción"]),
            invoiced_amount=row["Monto Facturado"],
            currency=ExcelRepository._optional_currency(row["Moneda"]),
            invoice_date=row["Fecha Factura"],
            verification_status=VerificationStatus(str(row["Estado Verificación"])),
            payment_status=PaymentStatus(str(row["Estado de Pago"])),
            source_hash=str(row["Hash Documento"]),
            source_files=ExcelRepository._split(row["Archivos Origen"]),
            uncertain_fields=ExcelRepository._split(row["Campos Inciertos"]),
            duplicate_of=ExcelRepository._split(row["Posible Duplicado De"]),
            registered_at=row["Fecha Registro"] or datetime.now(),
            observations=ExcelRepository._split(row["Acción / Observaciones"]),
        )


    @staticmethod
    def _payment_from_row(
        row: dict[str, object],
    ) -> PaymentRecord:
        return PaymentRecord(
            payment_id=str(row["ID Pago"]),
            quotation_code=str(row["Cód. Cotización"]),
            invoice_id=ExcelRepository._optional_str(
                row["ID Factura"]
            ),
            payment_type=DocumentType(
                str(row["Tipo de Pago"])
            ),
            payment_entity=ExcelRepository._optional_str(
                row["Entidad"]
            ),
            operation_number=ExcelRepository._optional_str(
                row["N° Operación"]
            ),
            paid_amount=row["Monto Pagado"],
            currency=ExcelRepository._optional_currency(
                row["Moneda"]
            ),
            payment_date=row["Fecha Pago"],
            verification_status=VerificationStatus(
                str(row["Estado Verificación"])
            ),
            reconciliation_status=PaymentLinkStatus(
                str(row["Estado Conciliación"])
            ),
            source_hash=str(row["Hash Documento"]),
            source_files=ExcelRepository._split(
                row["Archivos Origen"]
            ),
            duplicate_of=ExcelRepository._split(
                row["Posible Duplicado De"]
            ),
            registered_at=(
                row["Fecha Registro"] or datetime.now()
            ),
            observations=ExcelRepository._split(
                row["Acción / Observaciones"]
            ),
        )

    @staticmethod
    def _optional_str(value: object) -> str | None:
        return None if value in (None, "") else str(value)

    @staticmethod
    def _optional_currency(value: object) -> Currency | None:
        return None if value in (None, "") else Currency(str(value))

    @staticmethod
    def _optional_document_type(value: object) -> DocumentType | None:
        return None if value in (None, "") else DocumentType(str(value))

    @staticmethod
    def _split(value: object) -> list[str]:
        if value in (None, ""):
            return []
        return [part.strip() for part in str(value).split("|") if part.strip()]
