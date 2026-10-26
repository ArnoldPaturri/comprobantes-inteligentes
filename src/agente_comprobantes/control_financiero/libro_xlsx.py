"""Generación de XLSX válidos con openpyxl, sin tocar el Excel operativo.

Conserva los contratos HojaReporte y Formula, pero delega la escritura OOXML
a una librería mantenida en lugar de construir XML manualmente.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties

COLOR_PRIMARIO = "714B67"
COLOR_SECUNDARIO = "017E84"
COLOR_TEXTO = "243040"
COLOR_ALERTA = "FFF5DE"
FORMATO_IMPORTE = '#,##0.00;[Red](#,##0.00)'


@dataclass(frozen=True)
class Formula:
    """Una fórmula generada por código; resultado esperado para pruebas."""

    expresion: str
    resultado: Decimal


@dataclass
class HojaReporte:
    nombre: str
    titulo: str
    subtitulo: str
    encabezados: list[str]
    filas: list[list[object]]
    anchos: list[int]
    columnas_monetarias: set[int] = field(default_factory=set)


def _texto(valor: object) -> str:
    texto = "" if valor is None else str(valor)
    # Excel/OOXML no admite los caracteres de control indicados abajo.
    return "".join(
        letra for letra in texto if ord(letra) >= 32 or letra in "\t\n\r"
    )


def _validar(hojas: list[HojaReporte]) -> None:
    if not hojas:
        raise ValueError("El reporte debe incluir al menos una hoja.")
    nombres = [hoja.nombre for hoja in hojas]
    if len(set(nombres)) != len(nombres):
        raise ValueError("Los nombres de hojas deben ser únicos.")
    for hoja in hojas:
        if (
            not hoja.nombre
            or len(hoja.nombre) > 31
            or any(c in hoja.nombre for c in "[]:*?/\\")
        ):
            raise ValueError("El nombre de hoja no es válido.")
        cantidad = len(hoja.encabezados)
        if not cantidad or cantidad != len(hoja.anchos):
            raise ValueError("Los encabezados y anchos no coinciden.")
        if any(len(fila) != cantidad for fila in hoja.filas):
            raise ValueError("Las filas no coinciden con el número de columnas.")
        if any(n < 1 or n > cantidad for n in hoja.columnas_monetarias):
            raise ValueError("La columna monetaria está fuera del reporte.")


def _escribir_celda(celda, valor: object) -> None:
    if isinstance(valor, Formula):
        celda.value = "=" + valor.expresion.removeprefix("=")
    elif isinstance(valor, (Decimal, int)) and not isinstance(valor, bool):
        celda.value = valor
    elif valor is not None:
        # Los textos procedentes de OCR pueden empezar con =, +, - o @.
        # Forzarlos como cadenas impide que se ejecuten como fórmulas.
        celda.value = _texto(valor)
        celda.data_type = "s"


def _dibujar_hoja(libro: Workbook, hoja: HojaReporte) -> None:
    ws = libro.create_sheet(title=hoja.nombre)
    total = len(hoja.encabezados)
    ultima = get_column_letter(total)
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A6"

    ws.merge_cells(f"A1:{ultima}1")
    ws.merge_cells(f"A2:{ultima}2")
    ws["A1"] = _texto(hoja.titulo)
    ws["A2"] = _texto(hoja.subtitulo)
    ws["A1"].font = Font(name="Aptos", size=17, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor=COLOR_PRIMARIO)
    ws["A2"].font = Font(name="Aptos", size=10, color=COLOR_TEXTO)
    ws["A1"].alignment = Alignment(vertical="center", indent=1)
    ws["A2"].alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 34
    ws.row_dimensions[2].height = 25
    ws.row_dimensions[5].height = 32

    for i, (titulo, ancho) in enumerate(zip(hoja.encabezados, hoja.anchos, strict=True), 1):
        ws.column_dimensions[get_column_letter(i)].width = ancho
        celda = ws.cell(row=5, column=i, value=_texto(titulo))
        celda.data_type = "s"
        celda.fill = PatternFill("solid", fgColor=COLOR_SECUNDARIO)
        celda.font = Font(name="Aptos", bold=True, color="FFFFFF", size=10)
        celda.alignment = Alignment(
            vertical="center", wrap_text=True, horizontal="center"
        )

    for n, fila in enumerate(hoja.filas, 6):
        ws.row_dimensions[n].height = 23
        for i, valor in enumerate(fila, 1):
            celda = ws.cell(row=n, column=i)
            _escribir_celda(celda, valor)
            celda.font = Font(name="Aptos", size=10, color=COLOR_TEXTO)
            celda.alignment = Alignment(vertical="center", wrap_text=True)
            if i in hoja.columnas_monetarias:
                celda.number_format = FORMATO_IMPORTE
            if valor in ("OBSERVADA", "PENDIENTE"):
                celda.fill = PatternFill("solid", fgColor=COLOR_ALERTA)
        if n % 2 == 0:
            # Bandas suaves sin cubrir el formato de observaciones.
            for i in range(1, total + 1):
                celda = ws.cell(row=n, column=i)
                if celda.fill.patternType is None:
                    celda.fill = PatternFill("solid", fgColor="F7F7F9")

    # Los filtros y filas congeladas quedan definidos mediante openpyxl,
    # con el orden OOXML correcto para Microsoft Excel.
    ws.auto_filter.ref = f"A5:{ultima}{max(5, 5 + len(hoja.filas))}"
    ws.print_options.horizontalCentered = True
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0


def crear_archivo_xlsx(hojas: list[HojaReporte]) -> bytes:
    """Crea un Excel nuevo en memoria sin alterar ningún libro existente."""
    _validar(hojas)
    libro = Workbook()
    libro.remove(libro.active)
    libro.calculation = CalcProperties(
        calcMode="auto", fullCalcOnLoad=True, forceFullCalc=True
    )
    for hoja in hojas:
        _dibujar_hoja(libro, hoja)
    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()
