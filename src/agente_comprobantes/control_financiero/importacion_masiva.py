"""Lectura masiva conservadora de hojas Excel para control financiero.

No escribe en la fuente, no confirma pagos ni clasifica automáticamente
operaciones como fletes. La vista previa separa válidos y observaciones.
"""

from __future__ import annotations

import hashlib
import io
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from itertools import islice

from openpyxl import load_workbook

from agente_comprobantes.control_financiero.modelos import (
    DatosOrigenExcel,
    MovimientoBancario,
    OperacionFlete,
)
from agente_comprobantes.control_financiero.repositorio_json import (
    RepositorioControlFletes,
)
from agente_comprobantes.dominio.enumeraciones import Moneda

MAX_ARCHIVO = 8 * 1024 * 1024
MAX_FILAS = 5000

COLUMNAS = {
    "factura": ("numero factura", "n factura", "nro factura", "factura"),
    "proveedor": ("proveedor", "razon social"),
    "moneda": ("moneda", "divisa"),
    "facturado": ("monto facturado", "facturado", "importe factura", "total factura"),
    "identificador": ("id factura", "id expediente"),
    "ruc": ("ruc proveedor", "ruc"),
    "guia": ("nro de guia", "n guia", "numero guia", "nro guia", "guia"),
    "fecha": ("fecha factura", "fecha de factura"),
    "cotizacion": ("codigo cotizacion", "cod cotizacion", "cotizacion"),
    "presupuesto": ("monto cotizado", "cotizado", "presupuesto"),
    "pago": ("monto bancarizado", "pago declarado"),
    "referencia": ("nro bancarizacion", "numero bancarizacion", "nro de bancarizacion"),
    "detraccion": ("detraccion",),
    "nota_credito": ("nota de credito", "nota credito"),
    "estado": ("estado", "estado de pago", "estado pago"),
    "diferencia": ("diferencia", "diferencia bancaria", "diferencia bancaria factura"),
    "observaciones": (
        "accion observaciones", "accion observacion", "acciones observaciones",
        "observaciones", "accion requerida",
    ),
}
OBLIGATORIAS = {"factura", "proveedor", "moneda", "facturado"}
ALIAS = {alias: campo for campo, nombres in COLUMNAS.items() for alias in nombres}


def _normalizar(valor: object) -> str:
    if valor is None:
        return ""
    sin_tildes = "".join(
        letra for letra in unicodedata.normalize("NFKD", str(valor))
        if not unicodedata.combining(letra)
    )
    return re.sub(r"[^a-z0-9]+", " ", sin_tildes.casefold()).strip()


def _texto(valor: object) -> str:
    return "" if valor is None else str(valor).strip()


def _moneda(valor: object) -> Moneda:
    texto = _texto(valor)
    codigo = _normalizar(texto)
    monedas = {
        "pen": Moneda.PEN,
        "soles": Moneda.PEN,
        "s": Moneda.PEN,
        "usd": Moneda.USD,
        "us": Moneda.USD,
        "dolares": Moneda.USD,
        "eur": Moneda.EUR,
        "euros": Moneda.EUR,
    }
    if codigo in monedas:
        return monedas[codigo]

    # Formatos reales de hojas contables: PEN (S/.), USD (US$), etc.
    # El símbolo debe ser compatible con el código: nunca inferir una
    # moneda cuando la etiqueta es contradictoria o desconocida.
    etiqueta = re.fullmatch(
        r"(PEN|USD|EUR)\s*\(\s*([^()]*)\s*\)",
        texto,
        flags=re.IGNORECASE,
    )
    if etiqueta:
        iso, simbolo = etiqueta.groups()
        iso = iso.upper()
        simbolo = re.sub(r"\s+", "", simbolo).upper().rstrip(".")
        simbolos_validos = {
            "PEN": {"S/", "SOLES", "PEN"},
            "USD": {"US$", "$", "DOLARES", "USD"},
            "EUR": {"€", "EUROS", "EUR"},
        }
        if simbolo in simbolos_validos[iso]:
            return Moneda(iso)

    raise ValueError(f"Moneda desconocida: {texto or 'vacía'}")


def _importe(valor: object, *, permitir_cero: bool = False) -> Decimal:
    if isinstance(valor, bool) or valor is None:
        raise ValueError("Importe vacío o no numérico.")
    texto = _texto(valor).replace("S/", "").replace("US$", "").replace("$", "")
    texto = texto.replace(" ", "")
    if not texto:
        raise ValueError("Importe vacío.")
    if "," in texto and "." in texto:
        if texto.rfind(",") > texto.rfind("."):
            texto = texto.replace(".", "").replace(",", ".")
        else:
            texto = texto.replace(",", "")
    elif "," in texto:
        texto = (
            texto.replace(",", ".")
            if len(texto.rsplit(",", 1)[-1]) <= 2
            else texto.replace(",", "")
        )
    try:
        importe = Decimal(texto)
    except InvalidOperation as exc:
        raise ValueError(f"Importe inválido: {_texto(valor)}") from exc
    if (
        not importe.is_finite()
        or importe < 0
        or (importe == 0 and not permitir_cero)
        or importe > Decimal("999999999999.99")
        or importe != importe.quantize(Decimal("0.01"))
    ):
        raise ValueError("Importe fuera del rango monetario permitido.")
    return importe


def _fecha(valor: object) -> date | None:
    if valor is None or _texto(valor) == "":
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    texto = _texto(valor)
    for plantilla in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(texto, plantilla).date()
        except ValueError:
            continue
    raise ValueError(f"Fecha de factura no reconocida: {texto}")


def _clave(factura: str, proveedor: str) -> str:
    return f"{_normalizar(factura)}|{_normalizar(proveedor)}"


def _id_generado(clave: str) -> str:
    return "FLE-" + hashlib.sha256(clave.encode("utf-8")).hexdigest()[:12].upper()


@dataclass(frozen=True)
class FilaImportada:
    hoja: str
    fila: int
    estado: str
    identificador: str
    factura: str
    proveedor: str
    moneda: str
    monto: str
    detalles: tuple[str, ...]
    operacion: OperacionFlete | None = None
    valores_origen: dict[str, str] = field(default_factory=dict)


@dataclass(frozen=True)
class VistaPrevia:
    nombre_hoja: str
    encabezado: int
    filas: list[FilaImportada]
    huella_destino: str
    columnas_origen: tuple[str, ...] = ()
    columnas_no_reconocidas: tuple[str, ...] = ()

    @property
    def listas(self) -> list[OperacionFlete]:
        return [fila.operacion for fila in self.filas
                if fila.estado == "LISTA" and fila.operacion is not None]

    @property
    def errores(self) -> int:
        return sum(f.estado == "ERROR" for f in self.filas)

    @property
    def duplicadas(self) -> int:
        return sum(f.estado == "DUPLICADA" for f in self.filas)

    @property
    def omitidas(self) -> int:
        return sum(f.estado == "OMITIDA" for f in self.filas)


def _encabezados(valores: tuple[object, ...]) -> dict[str, int]:
    columnas: dict[str, int] = {}
    for i, valor in enumerate(valores):
        campo = ALIAS.get(_normalizar(valor))
        if campo is not None and campo not in columnas:
            columnas[campo] = i
    return columnas


def _detectar(libro) -> tuple[object, int, dict[str, int]]:
    mejor = None
    for ws in libro.worksheets:
        # En modo lectura, algunas hojas de Excel no declaran su dimensión
        # o declaran un rango menor al contenido real. Se lee el XML real
        # sin confiar en ws.max_row (que incluso puede ser None).
        ws.reset_dimensions()
        for pos, fila in enumerate(
            ws.iter_rows(min_row=1, max_row=25, values_only=True),
            start=1,
        ):
            columnas = _encabezados(fila)
            if OBLIGATORIAS.issubset(columnas):
                prioridad = int("flete" in _normalizar(ws.title))
                puntaje = (len(columnas), prioridad)
                if mejor is None or puntaje > mejor[0]:
                    mejor = (puntaje, ws, pos, columnas)
    if mejor is None:
        raise ValueError(
            "No se encontraron las columnas Factura, Proveedor, Moneda y Monto Facturado."
        )
    return mejor[1], mejor[2], mejor[3]


def _convertir(
    datos: dict[str, object],
    datos_origen: DatosOrigenExcel | None = None,
) -> tuple[OperacionFlete, tuple[str, ...]]:
    factura = _texto(datos.get("factura"))
    proveedor = _texto(datos.get("proveedor"))
    if not factura or not proveedor:
        raise ValueError("Factura y proveedor son obligatorios.")
    moneda = _moneda(datos.get("moneda"))
    monto = _importe(datos.get("facturado"))
    clave = _clave(factura, proveedor)
    identificador = _texto(datos.get("identificador")) or _id_generado(clave)
    advertencias = [
        "Importada sin clasificación de flete ni acreditación bancaria.",
    ]
    movimientos = []
    valor_pago = datos.get("pago")
    if valor_pago is not None and _texto(valor_pago):
        bancarizado = _importe(valor_pago, permitir_cero=True)
        if bancarizado > 0:
            movimientos = [MovimientoBancario(
                moneda=moneda,
                importe_aplicado=bancarizado,
                referencia=_texto(datos.get("referencia")) or None,
                verificada=False,
                documento_respaldo=None,
            )]
            advertencias.append(
                "Monto bancarizado importado solo como declaración, sin respaldo."
            )
    if _texto(datos.get("estado")):
        advertencias.append(
            "El estado del Excel se conservó como dato de origen, sin acreditar."
        )
    if _texto(datos.get("diferencia")):
        advertencias.append(
            "La diferencia del Excel se conservó como dato de origen; "
            "no se asumió su fórmula de cálculo."
        )
    if _texto(datos.get("observaciones")):
        advertencias.append(
            "Las acciones u observaciones del Excel se conservaron para revisión."
        )
    if _texto(datos.get("detraccion")):
        advertencias.append("Debe revisarse la aplicabilidad y constancia de detracción.")
    nota = _normalizar(datos.get("nota_credito"))
    if nota and nota not in {"no aplica", "no", "ninguna", "n a"}:
        advertencias.append("Nota de crédito mencionada; falta registrar su documento.")
    presupuesto = datos.get("presupuesto")
    operacion = OperacionFlete(
        identificador=identificador,
        numero_factura=factura,
        proveedor=proveedor,
        ruc_proveedor=_texto(datos.get("ruc")) or None,
        moneda=moneda,
        monto_facturado=monto,
        fecha_factura=_fecha(datos.get("fecha")),
        numero_guia=_texto(datos.get("guia")) or None,
        codigo_cotizacion=_texto(datos.get("cotizacion")) or None,
        monto_cotizado=(
            _importe(presupuesto, permitir_cero=True)
            if presupuesto is not None and _texto(presupuesto) else None
        ),
        es_flete=None,
        detraccion_aplica=None,
        conciliacion_cerrada=False,
        movimientos_bancarios=movimientos,
        datos_origen=datos_origen,
    )
    return operacion, tuple(advertencias)


def preparar_importacion(
    datos_xlsx: bytes,
    repositorio: RepositorioControlFletes,
    *,
    nombre_archivo: str | None = None,
) -> VistaPrevia:
    """Solo lectura. Ningún dato se incorpora sin confirmación posterior."""
    if not isinstance(datos_xlsx, bytes) or not 0 < len(datos_xlsx) <= MAX_ARCHIVO:
        raise ValueError("El Excel debe tener entre 1 byte y 8 MB.")
    if not datos_xlsx.startswith(b"PK"):
        raise ValueError("El archivo no parece ser un libro .xlsx válido.")
    try:
        libro = load_workbook(io.BytesIO(datos_xlsx), read_only=True, data_only=True,
                            keep_links=False)
    except (OSError, ValueError, KeyError) as exc:
        raise ValueError("No se pudo leer el archivo Excel.") from exc

    try:
        ws, encabezado, columnas = _detectar(libro)
        encabezados = next(
            ws.iter_rows(
                min_row=encabezado, max_row=encabezado, values_only=True
            )
        )
        nombres = tuple(
            _texto(valor) if _texto(valor) else f"Columna {indice}"
            for indice, valor in enumerate(encabezados, start=1)
        )
        sin_mapear = tuple(
            nombre for nombre in nombres
            if ALIAS.get(_normalizar(nombre)) is None
            and not nombre.startswith("Columna ")
        )
        huella_archivo = hashlib.sha256(datos_xlsx).hexdigest()
        existentes = repositorio.listar()
        ids = {op.identificador.strip().casefold() for op in existentes}
        claves = {_clave(op.numero_factura, op.proveedor or "") for op in existentes}
        vistos: set[str] = set()
        filas: list[FilaImportada] = []
        for numero, valores in enumerate(
            islice(
                ws.iter_rows(min_row=encabezado + 1, values_only=True),
                MAX_FILAS + 1,
            ),
            start=encabezado + 1,
        ):
            if numero - encabezado > MAX_FILAS:
                raise ValueError(f"El límite por importación es {MAX_FILAS} filas.")
            if not any(x is not None and _texto(x) for x in valores):
                continue
            campos = {
                campo: valores[i] if i < len(valores) else None
                for campo, i in columnas.items()
            }
            valores_origen = {
                (nombre if nombres.count(nombre) == 1
                 else f"{nombre} [{indice + 1}]"): _texto(
                    valores[indice] if indice < len(valores) else None
                )
                for indice, nombre in enumerate(nombres)
            }
            datos_origen = DatosOrigenExcel(
                archivo=nombre_archivo,
                huella_archivo=huella_archivo,
                hoja=ws.title,
                fila=numero,
                campos=valores_origen,
            )
            factura = _texto(campos.get("factura"))
            proveedor = _texto(campos.get("proveedor"))
            # Las últimas filas suelen contener notas de control, listas
            # desplegables o pies de tabla sin una operación económica.
            # Si existe un importe, se mantiene como ERROR y no se oculta.
            if (
                not factura
                and not proveedor
                and not _texto(campos.get("facturado"))
                and not _texto(campos.get("pago"))
            ):
                filas.append(FilaImportada(
                    ws.title, numero, "OMITIDA", "", "", "",
                    _texto(campos.get("moneda")), "",
                    ("Fila sin factura, proveedor ni importes; no se importa.",),
                    valores_origen=valores_origen,
                ))
                continue
            try:
                op, avisos = _convertir(campos, datos_origen)
            except (ValueError, TypeError) as exc:
                filas.append(FilaImportada(
                    ws.title, numero, "ERROR", "", factura, proveedor,
                    _texto(campos.get("moneda")), _texto(campos.get("facturado")),
                    (str(exc),),
                    valores_origen=valores_origen,
                ))
                continue
            clave = _clave(op.numero_factura, op.proveedor or "")
            if (
                clave in claves or clave in vistos
                or op.identificador.strip().casefold() in ids
            ):
                estado = "DUPLICADA"
                avisos = (*avisos, "La factura ya está registrada o repetida en el lote.")
            else:
                estado = "LISTA"
                vistos.add(clave)
                ids.add(op.identificador.strip().casefold())
            filas.append(FilaImportada(
                ws.title, numero, estado, op.identificador, op.numero_factura,
                op.proveedor or "", op.moneda.value, str(op.monto_facturado),
                avisos, op if estado == "LISTA" else None,
                valores_origen=valores_origen,
            ))
        return VistaPrevia(
            ws.title,
            encabezado,
            filas,
            repositorio.huella(),
            columnas_origen=nombres,
            columnas_no_reconocidas=sin_mapear,
        )
    finally:
        libro.close()


def confirmar_importacion(
    previa: VistaPrevia,
    repositorio: RepositorioControlFletes,
) -> int:
    """Registra el lote en una sola escritura, revalidando el destino."""
    from agente_comprobantes.control_financiero.auditoria_importacion import (
        auditar_importacion,
    )

    if not previa.listas:
        raise ValueError("No hay operaciones nuevas para importar.")
    auditoria = auditar_importacion(previa)
    if auditoria.bloquea_importacion:
        raise ValueError(
            "Hay referencias o facturas que necesitan agrupación y revisión "
            "manual antes de importar. Consulta los hallazgos de la vista previa."
        )
    repositorio.agregar_lote(previa.listas, huella_esperada=previa.huella_destino)
    return len(previa.listas)
