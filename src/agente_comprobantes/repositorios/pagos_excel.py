from __future__ import annotations

from openpyxl import load_workbook

from agente_comprobantes.dominio import (
    EstadoVinculacionPago,
    RegistroFactura,
    RegistroPago,
    ResultadoConciliacionPago,
)
from agente_comprobantes.excepciones import ErrorRepositorio
from agente_comprobantes.repositorios.esquemas_excel import PAYMENT_HEADERS


class OperacionesPagosExcel:
    def get_payment(self, payment_id: str) -> RegistroPago | None:
        normalized = payment_id.strip().upper()

        for row in self.read_sheet("PAGOS"):
            if str(row["ID Pago"]).strip().upper() == normalized:
                return self._payment_from_row(row)

        return None

    def list_pending_payments(
        self,
        quotation_code: str,
    ) -> list[RegistroPago]:
        normalized = quotation_code.strip().upper()
        pending: list[RegistroPago] = []

        for row in self.read_sheet("PAGOS"):
            row_quotation = str(
                row["Cód. Cotización"] or ""
            ).strip().upper()

            if row_quotation != normalized:
                continue

            if (
                str(row["Estado Conciliación"] or "")
                != EstadoVinculacionPago.PENDING_INVOICE.value
            ):
                continue

            pending.append(self._payment_from_row(row))

        return pending

    def add_payment(
        self,
        payment: RegistroPago,
        invoice: RegistroFactura | None = None,
        reconciliation: ResultadoConciliacionPago | None = None,
    ) -> None:
        self.initialize()
        workbook = load_workbook(self._path)
        self._append_mapping(
            workbook["PAGOS"],
            PAYMENT_HEADERS,
            {
                "ID Pago": payment.payment_id,
                "ID Factura": payment.invoice_id,
                "Tipo de Pago": payment.payment_type,
                "Entidad": payment.payment_entity,
                "N° Operación": payment.operation_number,
                "Monto Pagado": payment.paid_amount,
                "Moneda": payment.currency,
                "Fecha Pago": payment.payment_date,
                "Estado Verificación": payment.verification_status,
                "Hash Documento": payment.source_hash,
                "Fecha Registro": payment.registered_at,
                "Acción / Observaciones": payment.observations,
                "Cód. Cotización": payment.quotation_code,
                "Estado Conciliación": payment.reconciliation_status,
                "Archivos Origen": payment.source_files,
                "Posible Duplicado De": payment.duplicate_of,
            },
        )

        if (
            invoice is not None
            and reconciliation is not None
            and reconciliation.applied
        ):
            self._update_invoice_payment_status(
                workbook["FACTURAS"],
                invoice.invoice_id,
                reconciliation.resulting_payment_status,
            )
            self._update_comparison_payment_status(
                workbook["COMPARATIVA"],
                invoice.quotation_code,
                invoice.invoice_number,
                reconciliation.resulting_payment_status,
            )
        self._save_atomic(workbook)

    def link_payment(
        self,
        payment: RegistroPago,
        invoice: RegistroFactura,
        reconciliation: ResultadoConciliacionPago,
    ) -> None:
        self.initialize()
        workbook = load_workbook(self._path)
        sheet = workbook["PAGOS"]

        headers = {
            cell.value: cell.column
            for cell in sheet[1]
        }

        for row_number in range(2, sheet.max_row + 1):
            current = sheet.cell(
                row=row_number,
                column=headers["ID Pago"],
            ).value

            if (
                str(current).strip().upper()
                != payment.payment_id.strip().upper()
            ):
                continue

            sheet.cell(
                row=row_number,
                column=headers["ID Factura"],
            ).value = payment.invoice_id

            sheet.cell(
                row=row_number,
                column=headers["Estado Conciliación"],
            ).value = payment.reconciliation_status.value

            sheet.cell(
                row=row_number,
                column=headers["Acción / Observaciones"],
            ).value = self._excel_value(payment.observations)

            break
        else:
            workbook.close()
            raise ErrorRepositorio(
                f"No existe el pago {payment.payment_id}."
            )

        if reconciliation.applied:
            self._update_invoice_payment_status(
                workbook["FACTURAS"],
                invoice.invoice_id,
                reconciliation.resulting_payment_status,
            )

            self._update_comparison_payment_status(
                workbook["COMPARATIVA"],
                invoice.quotation_code,
                invoice.invoice_number,
                reconciliation.resulting_payment_status,
            )

        self._save_atomic(workbook)
