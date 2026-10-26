from __future__ import annotations

import os
import tempfile
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from agente_comprobantes.dominio import (
    Moneda,
    TipoDocumento,
    RegistroFactura,
    EstadoVinculacionPago,
    RegistroPago,
    EstadoPago,
    RegistroCotizacion,
    EstadoCotizacion,
    EstadoVerificacion,
)
from agente_comprobantes.excepciones import ErrorRepositorio
from agente_comprobantes.repositorios.esquemas_excel import ILLEGAL_EXCEL_CHARACTER_PATTERN


class UtilidadesExcel:
    def read_sheet(self, sheet_name: str) -> list[dict[str, object]]:
        if not self._path.exists():
            return []
        workbook = load_workbook(self._path, data_only=True, read_only=True)
        try:
            if sheet_name not in workbook.sheetnames:
                raise ErrorRepositorio(f"No existe la hoja {sheet_name}.")
            sheet = workbook[sheet_name]
            rows = sheet.iter_rows(values_only=True)
            headers = next(rows, None)
            if not headers:
                return []
            return [dict(zip(headers, row, strict=False)) for row in rows if any(row)]
        finally:
            workbook.close()

    @staticmethod
    def _initialize_sheet(sheet: Any, headers: list[str]) -> None:
        sheet.append(headers)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
        fill = PatternFill("solid", fgColor="1F4E78")
        for cell in sheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for index, header in enumerate(headers, start=1):
            width = min(max(len(header) + 3, 14), 32)
            sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.row_dimensions[1].height = 28

    @staticmethod
    def _ensure_headers(sheet: Any, expected: list[str]) -> bool:
        actual = [cell.value for cell in sheet[1]]

        while actual and actual[-1] is None:
            actual.pop()

        if actual == expected:
            return False

        if actual != expected[: len(actual)]:
            raise ErrorRepositorio(
                f"La hoja {sheet.title} existe, pero sus columnas "
                "no coinciden con el contrato."
            )

        fill = PatternFill("solid", fgColor="1F4E78")

        for index, header in enumerate(
            expected[len(actual):],
            start=len(actual) + 1,
        ):
            cell = sheet.cell(row=1, column=index, value=header)
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = fill
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )
            sheet.column_dimensions[get_column_letter(index)].width = min(
                max(len(header) + 3, 14),
                32,
            )

        sheet.auto_filter.ref = (
            f"A1:{get_column_letter(len(expected))}{sheet.max_row}"
        )
        return True

    @classmethod
    def _append_mapping(cls, sheet: Any, headers: list[str], mapping: dict[str, Any]) -> None:
        sheet.append([cls._excel_value(mapping.get(header)) for header in headers])
        row_number = sheet.max_row
        for cell in sheet[row_number]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        cls._apply_number_formats(sheet, row_number, headers)
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{sheet.max_row}"

    @staticmethod
    def _apply_number_formats(sheet: Any, row_number: int, headers: list[str]) -> None:
        for index, header in enumerate(headers, start=1):
            cell = sheet.cell(row=row_number, column=index)
            if header.startswith("Monto") or header == "Diferencia ($)":
                cell.number_format = '#,##0.00;[Red]-#,##0.00'
            elif header == "% Desvío":
                cell.number_format = '0.0000;[Red]-0.0000'
            elif header.startswith("Fecha") and isinstance(cell.value, datetime):
                cell.number_format = "dd/mm/yyyy hh:mm"
            elif header.startswith("Fecha") and isinstance(cell.value, date):
                cell.number_format = "dd/mm/yyyy"

    @staticmethod
    def _excel_value(value: Any) -> Any:
        if isinstance(value, Decimal):
            return float(value)

        if hasattr(value, "value"):
            value = value.value

        if isinstance(value, list):
            value = " | ".join(
                str(item)
                for item in value
            )

        if isinstance(value, str):
            return ILLEGAL_EXCEL_CHARACTER_PATTERN.sub(
                "",
                value,
            )

        return value

    @staticmethod
    def _update_quotation_status(sheet: Any, code: str, status: EstadoCotizacion) -> None:
        headers = {cell.value: cell.column for cell in sheet[1]}
        code_col = headers["Cód. Cotización"]
        status_col = headers["Estado"]
        for row_number in range(2, sheet.max_row + 1):
            current = sheet.cell(row=row_number, column=code_col).value
            if str(current).strip().upper() == code.strip().upper():
                sheet.cell(row=row_number, column=status_col).value = status.value
                return

    @staticmethod
    def _update_invoice_payment_status(
        sheet: Any,
        invoice_id: str,
        status: EstadoPago,
    ) -> None:
        headers = {cell.value: cell.column for cell in sheet[1]}
        id_col = headers["ID Factura"]
        status_col = headers["Estado de Pago"]
        for row_number in range(2, sheet.max_row + 1):
            current = sheet.cell(row=row_number, column=id_col).value
            if str(current).strip().upper() == invoice_id.strip().upper():
                sheet.cell(row=row_number, column=status_col).value = status.value
                return

    @staticmethod
    def _update_comparison_payment_status(
        sheet: Any,
        quotation_code: str,
        invoice_number: str | None,
        status: EstadoPago,
    ) -> None:
        headers = {cell.value: cell.column for cell in sheet[1]}
        quotation_col = headers["Cód. Cotización"]
        invoice_col = headers["N° Factura"]
        status_col = headers["Estado de Pago"]
        for row_number in range(2, sheet.max_row + 1):
            current_quotation = sheet.cell(row=row_number, column=quotation_col).value
            current_invoice = sheet.cell(row=row_number, column=invoice_col).value
            if (
                str(current_quotation).strip().upper() == quotation_code.strip().upper()
                and str(current_invoice or "").strip().upper()
                == str(invoice_number or "").strip().upper()
            ):
                sheet.cell(row=row_number, column=status_col).value = status.value
                return

    def _save_atomic(self, workbook: Any) -> None:
        self._path.parent.mkdir(parents=True, exist_ok=True)
        temporary_path: Path | None = None
        try:
            with tempfile.NamedTemporaryFile(
                dir=self._path.parent, prefix=f".{self._path.stem}-", suffix=".xlsx", delete=False
            ) as temporary:
                temporary_path = Path(temporary.name)
            workbook.save(temporary_path)
            workbook.close()
            os.replace(temporary_path, self._path)
        except Exception as exc:
            workbook.close()
            if temporary_path and temporary_path.exists():
                temporary_path.unlink()
            raise ErrorRepositorio(f"No se pudo guardar el Excel: {self._path}") from exc

    @staticmethod
    def _quotation_from_row(row: dict[str, object]) -> RegistroCotizacion:
        return RegistroCotizacion(
            quotation_code=str(row["Cód. Cotización"]),
            provider=UtilidadesExcel._optional_str(row["Proveedor"]),
            provider_tax_id=UtilidadesExcel._optional_str(row["RUC Proveedor"]),
            concept=UtilidadesExcel._optional_str(row["Concepto / Descripción"]),
            quoted_amount=row["Monto Cotizado"],
            currency=UtilidadesExcel._optional_currency(row["Moneda"]),
            quotation_date=row["Fecha Cotización"],
            status=EstadoCotizacion(str(row["Estado"])),
            verification_status=EstadoVerificacion(str(row["Estado Verificación"])),
            source_hash=str(row["Hash Documento"]),
            source_files=UtilidadesExcel._split(row["Archivos Origen"]),
            uncertain_fields=UtilidadesExcel._split(row["Campos Inciertos"]),
            registered_at=row["Fecha Registro"] or datetime.now(),
            observations=UtilidadesExcel._split(row["Acción / Observaciones"]),
        )

    @staticmethod
    def _invoice_from_row(row: dict[str, object]) -> RegistroFactura:
        return RegistroFactura(
            invoice_id=str(row["ID Factura"]),
            quotation_code=str(row["Cód. Cotización"]),
            provider=UtilidadesExcel._optional_str(row["Proveedor"]),
            provider_tax_id=UtilidadesExcel._optional_str(row["RUC Proveedor"]),
            invoice_number=UtilidadesExcel._optional_str(row["N° Factura"]),
            concept=UtilidadesExcel._optional_str(row["Concepto / Descripción"]),
            invoiced_amount=row["Monto Facturado"],
            currency=UtilidadesExcel._optional_currency(row["Moneda"]),
            invoice_date=row["Fecha Factura"],
            verification_status=EstadoVerificacion(str(row["Estado Verificación"])),
            payment_status=EstadoPago(str(row["Estado de Pago"])),
            source_hash=str(row["Hash Documento"]),
            source_files=UtilidadesExcel._split(row["Archivos Origen"]),
            uncertain_fields=UtilidadesExcel._split(row["Campos Inciertos"]),
            duplicate_of=UtilidadesExcel._split(row["Posible Duplicado De"]),
            registered_at=row["Fecha Registro"] or datetime.now(),
            observations=UtilidadesExcel._split(row["Acción / Observaciones"]),
        )

    @staticmethod
    def _payment_from_row(
        row: dict[str, object],
    ) -> RegistroPago:
        return RegistroPago(
            payment_id=str(row["ID Pago"]),
            quotation_code=str(row["Cód. Cotización"]),
            invoice_id=UtilidadesExcel._optional_str(
                row["ID Factura"]
            ),
            payment_type=TipoDocumento(
                str(row["Tipo de Pago"])
            ),
            payment_entity=UtilidadesExcel._optional_str(
                row["Entidad"]
            ),
            operation_number=UtilidadesExcel._optional_str(
                row["N° Operación"]
            ),
            paid_amount=row["Monto Pagado"],
            currency=UtilidadesExcel._optional_currency(
                row["Moneda"]
            ),
            payment_date=row["Fecha Pago"],
            verification_status=EstadoVerificacion(
                str(row["Estado Verificación"])
            ),
            reconciliation_status=EstadoVinculacionPago(
                str(row["Estado Conciliación"])
            ),
            source_hash=str(row["Hash Documento"]),
            source_files=UtilidadesExcel._split(
                row["Archivos Origen"]
            ),
            duplicate_of=UtilidadesExcel._split(
                row["Posible Duplicado De"]
            ),
            registered_at=(
                row["Fecha Registro"] or datetime.now()
            ),
            observations=UtilidadesExcel._split(
                row["Acción / Observaciones"]
            ),
        )

    @staticmethod
    def _optional_str(value: object) -> str | None:
        return None if value in (None, "") else str(value)

    @staticmethod
    def _optional_currency(value: object) -> Moneda | None:
        return None if value in (None, "") else Moneda(str(value))

    @staticmethod
    def _optional_document_type(value: object) -> TipoDocumento | None:
        return None if value in (None, "") else TipoDocumento(str(value))

    @staticmethod
    def _split(value: object) -> list[str]:
        if value in (None, ""):
            return []
        return [part.strip() for part in str(value).split("|") if part.strip()]
