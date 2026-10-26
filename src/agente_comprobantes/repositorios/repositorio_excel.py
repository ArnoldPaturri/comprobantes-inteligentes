from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook

from agente_comprobantes.repositorios.base import RepositorioDocumentos
from agente_comprobantes.repositorios.cotizaciones_excel import OperacionesCotizacionesExcel
from agente_comprobantes.repositorios.esquemas_excel import SHEETS
from agente_comprobantes.repositorios.facturas_excel import OperacionesFacturasExcel
from agente_comprobantes.repositorios.pagos_excel import OperacionesPagosExcel
from agente_comprobantes.repositorios.referencias_excel import ReferenciasExcel
from agente_comprobantes.repositorios.utilidades_excel import UtilidadesExcel


class RepositorioExcel(
    OperacionesCotizacionesExcel,
    OperacionesFacturasExcel,
    OperacionesPagosExcel,
    ReferenciasExcel,
    UtilidadesExcel,
    RepositorioDocumentos,
):
    """Repositorio Excel compuesto por operaciones y utilidades especializadas."""

    def __init__(self, path: str | Path) -> None:
        self._path = Path(path).expanduser().resolve()

    @property
    def ubicacion(self) -> str:
        return str(self._path)

    def inicializar(self) -> None:
        self._path.parent.mkdir(parents=True, exist_ok=True)
        if self._path.exists():
            workbook = load_workbook(self._path)
            changed = False
            for name, headers in SHEETS.items():
                if name not in workbook.sheetnames:
                    sheet = workbook.create_sheet(name)
                    self._inicializar_hoja(sheet, headers)
                    changed = True
                else:
                    changed = self._verificar_encabezados(workbook[name], headers) or changed
            if changed:
                self._guardar_atomico(workbook)
            else:
                workbook.close()
            return

        workbook = Workbook()
        default = workbook.active
        workbook.remove(default)
        for name, headers in SHEETS.items():
            sheet = workbook.create_sheet(name)
            self._inicializar_hoja(sheet, headers)
        self._guardar_atomico(workbook)
